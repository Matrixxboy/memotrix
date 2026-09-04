from pathlib import Path
from typing import Any, Dict, List

from .base import BaseExtractor
from .media import describe_image_asset, media_dir_for, save_image_bytes
from memotrix.utils.outputSturcture import build_document


class DOCXExtractor(BaseExtractor):
    supported_extensions = (".docx",)

    def extract(self, path: Path, *, describe_images: bool = True):
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover
            from memotrix.utils.exceptions import MissingDependencyError
            raise MissingDependencyError("python-docx is required for DOCX extraction. Run `pip install memotrix[docx]`.") from exc

        path = Path(path)
        doc = DocxDocument(str(path))
        out_dir = media_dir_for(path)
        skip_ai = not describe_images

        paragraphs_and_headings: List[str] = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                continue
            style_name = getattr(p.style, "name", "") or ""
            if "Heading 1" in style_name:
                paragraphs_and_headings.append(f"# {txt}")
            elif "Heading 2" in style_name:
                paragraphs_and_headings.append(f"## {txt}")
            elif "Heading" in style_name:
                paragraphs_and_headings.append(f"### {txt}")
            else:
                paragraphs_and_headings.append(txt)

        table_blocks: List[str] = []
        tables: List[Dict[str, Any]] = []
        for table_index, table in enumerate(doc.tables):
            rows: List[List[str]] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(cells)
            if not rows:
                continue
            headers = rows[0]
            data_rows = rows[1:] if len(rows) > 1 else []
            header_line = " | ".join(headers)
            lines = [f"Table {table_index + 1} columns: {header_line}"]
            for row in data_rows:
                width = len(headers)
                padded = (row + [""] * width)[:width]
                lines.append(" | ".join(padded))
            block = "\n".join(lines)
            table_blocks.append(block)
            tables.append(
                {
                    "headers": headers,
                    "rows": data_rows,
                    "text": block,
                }
            )

        text = "\n\n".join(paragraphs_and_headings + table_blocks)

        images: List[Dict[str, Any]] = []
        image_index = 0
        for rel in doc.part.rels.values():
            reltype = getattr(rel, "reltype", "") or ""
            if "image" not in reltype:
                continue
            try:
                blob = rel.target_part.blob
            except Exception as e:
                from memotrix.utils.trace import mlog
                mlog("docx", f"Warning: Failed to extract image blob from rel {reltype}: {e}")
                continue
            if not blob:
                continue

            image_index += 1
            img_path = save_image_bytes(
                blob,
                out_dir,
                stem=f"docx_img{image_index}",
            )
            images.append(
                describe_image_asset(
                    img_path,
                    context=f"DOCX embedded image {image_index}",
                    skip_ai=skip_ai,
                )
            )

        # If the doc is essentially image-only, fold captions into body text
        if len(text) < 40 and images:
            from .media import image_descriptions_as_text

            caption_text = image_descriptions_as_text(images)
            text = f"{text}\n\n{caption_text}".strip() if text else caption_text

        return build_document(
            path,
            text,
            images=images,
            tables=tables,
            extra_metadata={
                "file_type": "docx",
                "library": "python-docx",
                "extracted_image_count": len(images),
            },
        )


