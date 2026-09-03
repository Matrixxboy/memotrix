from pathlib import Path
from typing import Any, Dict, List

from .base import BaseExtractor
from .media import describe_image_asset, image_descriptions_as_text, media_dir_for, save_image_bytes
from memotrix.utils.outputSturcture import build_document


class PPTXExtractor(BaseExtractor):
    supported_extensions = (".pptx",)

    def extract(self, path: Path, *, describe_images: bool = True):
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-pptx is required for PPTX extraction") from exc

        from memotrix.utils.trace import mlog

        path = Path(path)
        mlog("pptx", f"extract start: {path.name}")
        presentation = Presentation(str(path))
        out_dir = media_dir_for(path)
        skip_ai = not describe_images

        slides: List[str] = []
        images: List[Dict[str, Any]] = []
        image_index = 0

        for slide_index, slide in enumerate(presentation.slides):
            slide_no = slide_index + 1
            parts: List[str] = []

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text and shape.text.strip():
                    parts.append(shape.text.strip())

                # Embedded pictures (including those sitting between text)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image
                        blob = image.blob
                    except Exception:
                        continue
                    if not blob:
                        continue
                    image_index += 1
                    img_path = save_image_bytes(
                        blob,
                        out_dir,
                        stem=f"slide{slide_no}_img{image_index}",
                    )
                    asset = describe_image_asset(
                        img_path,
                        context=f"PPTX slide {slide_no} image {image_index}",
                        skip_ai=skip_ai,
                    )
                    images.append(asset)
                    desc = (asset.get("description") or "").strip()
                    if desc:
                        parts.append(f"[Image] {desc}")

            if parts:
                slides.append(f"## Slide {slide_no}\n\n" + "\n".join(parts))
            elif any(
                (img.get("context") or "").startswith(f"PPTX slide {slide_no}")
                for img in images
            ):
                # Image-only slide — captions already in images[]; add a stub section
                slide_caps = [
                    img["description"]
                    for img in images
                    if (img.get("context") or "").startswith(f"PPTX slide {slide_no}")
                    and img.get("description")
                ]
                if slide_caps:
                    slides.append(
                        f"## Slide {slide_no} (image)\n\n" + "\n\n".join(slide_caps)
                    )

        text = "\n\n".join(slides)
        if len(text) < 40 and images:
            caption_text = image_descriptions_as_text(images)
            text = f"{text}\n\n{caption_text}".strip() if text else caption_text

        mlog(
            "pptx",
            f"extract done: {path.name} slides={len(presentation.slides)} "
            f"images={len(images)} text_chars={len(text)}",
        )
        return build_document(
            path,
            text,
            images=images,
            extra_metadata={
                "file_type": "pptx",
                "library": "python-pptx",
                "slide_count": len(presentation.slides),
                "extracted_image_count": len(images),
            },
        )
