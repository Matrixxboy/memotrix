"""Domain extractors: graphs, geo, FHIR, epub, email, chat, logs."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest

from memotrix import Memory
from memotrix.config import IngestConfig, MemoryConfig
from memotrix.embeddings import FakeEmbeddings
from memotrix.filetypes import extract_file, supported_extensions


def test_supported_extensions_include_domain_types():
    ext = set(supported_extensions())
    for item in (".nt", ".graphml", ".geojson", ".eml", ".mbox", ".log", ".epub", ".jsonl"):
        assert item in ext


def test_ntriples_knowledge_graph(tmp_path: Path):
    path = tmp_path / "org.nt"
    path.write_text(
        "<http://ex.org/Alice> <http://ex.org/works_at> <http://ex.org/Acme> .\n"
        "<http://ex.org/Vessel> <http://ex.org/located_in> <http://ex.org/Port_of_Rotterdam> .\n",
        encoding="utf-8",
    )
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "knowledge_graph"
    assert "Alice works_at Acme" in doc.text
    assert "Vessel located_in Port_of_Rotterdam" in doc.text


def test_graphml_knowledge_graph(tmp_path: Path):
    path = tmp_path / "org.graphml"
    path.write_text(
        """<?xml version="1.0"?>
<graphml>
  <graph>
    <node id="a"><data key="label">Alice</data></node>
    <node id="b"><data key="label">Acme</data></node>
    <edge source="a" target="b"><data key="label">works_at</data></edge>
  </graph>
</graphml>
""",
        encoding="utf-8",
    )
    doc = extract_file(path)
    assert "Alice works_at Acme" in doc.text


def test_geojson_feature(tmp_path: Path):
    path = tmp_path / "ports.geojson"
    path.write_text(
        json.dumps(
            {
                "type": "FeatureCollection",
                "features": [
                    {
                        "type": "Feature",
                        "properties": {"name": "Port of Singapore", "country": "SG"},
                        "geometry": {"type": "Point", "coordinates": [103.82, 1.26]},
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "geojson"
    assert "Port of Singapore" in doc.text
    assert "103.82" in doc.text


def test_json_sniff_geojson_not_generic_dump(tmp_path: Path):
    path = tmp_path / "place.json"
    path.write_text(
        json.dumps(
            {
                "type": "Feature",
                "properties": {"name": "EEZ boundary"},
                "geometry": {"type": "Point", "coordinates": [0.0, 0.0]},
            }
        ),
        encoding="utf-8",
    )
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "geojson"
    assert "EEZ boundary" in doc.text


def test_generic_json_still_pretty_printed(tmp_path: Path):
    path = tmp_path / "config.json"
    path.write_text(json.dumps({"foo": 1, "bar": "ok"}), encoding="utf-8")
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "json"
    assert '"foo"' in doc.text


def test_fhir_patient_and_bundle(tmp_path: Path):
    path = tmp_path / "patient.json"
    path.write_text(
        json.dumps(
            {
                "resourceType": "Bundle",
                "entry": [
                    {
                        "resource": {
                            "resourceType": "Patient",
                            "id": "p1",
                            "name": [{"given": ["Ada"], "family": "Lovelace"}],
                            "gender": "female",
                            "birthDate": "1815-12-10",
                        }
                    },
                    {
                        "resource": {
                            "resourceType": "Condition",
                            "id": "c1",
                            "code": {"text": "hypertension"},
                        }
                    },
                ],
            }
        ),
        encoding="utf-8",
    )
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "fhir"
    assert "Ada Lovelace" in doc.text
    assert "hypertension" in doc.text


def test_email_eml(tmp_path: Path):
    path = tmp_path / "note.eml"
    path.write_text(
        "From: ops@port.test\n"
        "To: captain@vessel.test\n"
        "Subject: Port delay\n"
        "Date: Thu, 3 Sep 2026 10:00:00 +0000\n"
        "\n"
        "Vessel delayed at Singapore due to weather.\n",
        encoding="utf-8",
    )
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "email"
    assert "Port delay" in doc.text
    assert "Singapore" in doc.text


def test_chat_json_and_whatsapp(tmp_path: Path):
    chat_json = tmp_path / "room.json"
    chat_json.write_text(
        json.dumps(
            [
                {"role": "user", "content": "Where is the vessel?"},
                {"role": "assistant", "content": "Alongside berth 12."},
            ]
        ),
        encoding="utf-8",
    )
    doc = extract_file(chat_json)
    assert doc.metadata["file_type"] == "chat"
    assert "berth 12" in doc.text

    wa = tmp_path / "whatsapp.txt"
    wa.write_text(
        "[01/01/2026, 10:00:00] Alice: Lesson starts at 9am\n"
        "[01/01/2026, 10:01:00] Bob: Bring the lab notes\n",
        encoding="utf-8",
    )
    doc2 = extract_file(wa)
    assert doc2.metadata["file_type"] == "chat"
    assert "lab notes" in doc2.text


def test_plain_txt_is_not_chat(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("This is a lecture outline about photosynthesis.\n", encoding="utf-8")
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "text"


def test_log_windows(tmp_path: Path):
    path = tmp_path / "app.log"
    lines = [f"2026-09-03T10:00:{i:02d}Z engine temperature {i}" for i in range(12)]
    path.write_text("\n".join(lines), encoding="utf-8")
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "log"
    assert "engine temperature" in doc.text
    assert "2026-09-03T10:00:00Z" in doc.text


def test_epub_chapter_text(tmp_path: Path):
    path = tmp_path / "book.epub"
    container = """<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>
"""
    opf = """<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" unique-identifier="id" version="3.0">
  <manifest>
    <item id="ch1" href="ch1.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine>
    <itemref idref="ch1"/>
  </spine>
</package>
"""
    html = """<html xmlns="http://www.w3.org/1999/xhtml"><body>
<p>Photosynthesis converts light into chemical energy.</p>
</body></html>"""
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr("META-INF/container.xml", container)
        zf.writestr("OEBPS/content.opf", opf)
        zf.writestr("OEBPS/ch1.xhtml", html)
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "epub"
    assert "Photosynthesis" in doc.text


def test_domain_files_are_searchable(tmp_path: Path):
    path = tmp_path / "facts.nt"
    path.write_text(
        "<http://ex.org/DrugX> <http://ex.org/treats> <http://ex.org/Hypertension> .\n",
        encoding="utf-8",
    )
    memory = Memory(
        embeddings=FakeEmbeddings(dim=8),
        backend="memory",
        config=MemoryConfig(ingest=IngestConfig(enable_dedup=False, describe_images=False)),
    )
    memory.add(path)
    hits = memory.search("what does DrugX treat?", top_k=3)
    assert hits
    blob = " ".join(h.get("chunk_text") or "" for h in hits)
    assert "Hypertension" in blob
    memory.close()


def test_catalog_json_is_not_chat(tmp_path: Path):
    path = tmp_path / "catalog.json"
    path.write_text(
        json.dumps([{"name": "Widget", "text": "A tool for ports"}]),
        encoding="utf-8",
    )
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "json"
    assert "Widget" in doc.text


def test_json_bom_extracts(tmp_path: Path):
    path = tmp_path / "bom.json"
    path.write_bytes(b"\xef\xbb\xbf" + json.dumps({"hello": "bom"}).encode("utf-8"))
    doc = extract_file(path)
    assert doc.metadata["file_type"] == "json"
    assert "hello" in doc.text


def test_generic_zip_is_unsupported(tmp_path: Path):
    from memotrix.utils.exceptions import UnsupportedDocumentTypeError

    path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("readme.txt", "not a scorm package")
    try:
        extract_file(path)
        raise AssertionError("expected UnsupportedDocumentTypeError")
    except UnsupportedDocumentTypeError as exc:
        assert "SCORM" in str(exc)


def test_xls_is_unsupported(tmp_path: Path):
    from memotrix.utils.exceptions import UnsupportedDocumentTypeError

    path = tmp_path / "legacy.xls"
    path.write_bytes(b"not a real workbook")
    try:
        extract_file(path)
        raise AssertionError("expected UnsupportedDocumentTypeError")
    except UnsupportedDocumentTypeError as exc:
        assert ".xlsx" in str(exc)


def test_docx_tables_and_skip_vision(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document as DocxDocument

    path = tmp_path / "table.docx"
    document = DocxDocument()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "42 million"
    document.save(path)
    doc = extract_file(path, describe_images=False)
    blob = (doc.text or "") + " ".join(t.get("text") or "" for t in doc.tables)
    assert "42 million" in blob
    assert doc.tables

    png = tmp_path / "tiny.png"
    png.write_bytes(
        bytes.fromhex(
            "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
            "0000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
        )
    )
    pictured = tmp_path / "pictured.docx"
    pictured_doc = DocxDocument()
    pictured_doc.add_picture(str(png))
    pictured_doc.save(pictured)
    with patch(
        "memotrix.utils.ai_integration.describe_image",
        side_effect=AssertionError("vision should not run"),
    ):
        extract_file(pictured, describe_images=False)
